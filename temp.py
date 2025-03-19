import pandas as pd
import sys
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

# Define the paths for input and output files
input_file = sys.argv[1]
output_file = input_file[:input_file.rfind('.')] + '.xlsx'

# Initialize lists to store the parsed data
timestamps = []
temperatures = []

def parse_line(timestamp_str, temperature_str):
    """
    Parses a timestamp and temperature line based on the format detected.
    """
    try:
        timestamp = datetime.strptime(timestamp_str, '%a %b %d %H:%M:%S UTC %Y')
    except ValueError:
        try:
            timestamp = datetime.strptime(timestamp_str, '%a %d %b %Y %I:%M:%S %p UTC')
        except ValueError:
            print(f"Error parsing timestamp: {timestamp_str}")
            return None, None
    
    # Parse temperature safely
    if '=' not in temperature_str:
        print(f"Skipping malformed temperature data: {temperature_str}")
        return timestamp, None  # Skip if no valid temperature

    try:
        temperature = float(temperature_str.split('=')[1].replace("'C", ""))
    except ValueError:
        print(f"Skipping invalid temperature data: {temperature_str}")
        return timestamp, None

    return timestamp, temperature

# Process file
with open(input_file, 'r') as file:
    lines = file.readlines()
    for i in range(0, len(lines), 2):  # Process in pairs (timestamp, temperature)
        if i + 1 >= len(lines):
            print(f"Skipping last unpaired line: {lines[i]}")
            break  # Avoid accessing an out-of-bounds index
        
        timestamp_str = lines[i].strip()
        temperature_str = lines[i + 1].strip()
        timestamp, temperature = parse_line(timestamp_str, temperature_str)
        
        if timestamp is not None and temperature is not None:
            timestamps.append(timestamp)
            temperatures.append(temperature)

# Create a DataFrame
df = pd.DataFrame({
    'Date/Time': timestamps,
    'Temperature (°C)': temperatures
})

# Save the DataFrame to an Excel file
df.to_excel(output_file, index=False)

# Load the workbook and select the active worksheet
wb = load_workbook(output_file)
ws = wb.active

# Create a line chart
chart = LineChart()
chart.title = "Temperature Over Time"
chart.style = 13
chart.y_axis.title = "Temperature (°C)"
chart.x_axis.title = "Date/Time"

# Set the data range for the chart
data = Reference(ws, min_col=2, min_row=2, max_row=len(df) + 1)  # Temperature column
dates = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)  # Date/Time column
chart.add_data(data, titles_from_data=True)
chart.set_categories(dates)

# Position the chart on the worksheet
ws.add_chart(chart, "D5")

# Save the workbook with the chart
wb.save(output_file)
print(f"Data and chart have been successfully written to {output_file}")
